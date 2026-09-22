# Builds example outputs (JSON, CSV, Excel) from a parsed price list, 3-5 services per delivery mode.
# Usage: python scripts/make-examples.py .tmp/sk.json examples
import csv, json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PICK = {
    "HOME": ["home:SK:131", "home:DE:13613", "home:BE:4832", "home:HU:3828", "home:IE:24810"],
    "PACKETA_PICKUP_POINT": ["packeta_pickup_point:SK:packeta-z-point-pp", "packeta_pickup_point:CZ:packeta-z-point-pp", "packeta_pickup_point:HU:packeta-z-point-pp"],
    "PACKETA_BOX": ["packeta_box:SK:packeta-z-box-pp", "packeta_box:CZ:packeta-z-box-pp", "packeta_box:RO:packeta-z-box-pp"],
    "PARTNER_PICKUP_POINT": ["partner_pickup_point:BE:7910", "partner_pickup_point:FR:12889", "partner_pickup_point:HR:10619", "partner_pickup_point:DE:6828"],
    "PARTNER_BOX": ["partner_box:BG:26067", "partner_box:PL:3060", "partner_box:HU:32970", "partner_box:GR:20409"],
}

src, out = Path(sys.argv[1]), Path(sys.argv[2])
out.mkdir(parents=True, exist_ok=True)
pl = json.loads(src.read_text(encoding="utf-8"))
by_id = {s["id"]: s for s in pl["services"]}
picked = [by_id[i] for ids in PICK.values() for i in ids]

example = {**pl, "services": picked, "returns": {**pl["returns"], "prices": pl["returns"]["prices"][:5]}}
example.pop("warnings", None)
(out / "sk-example.json").write_text(json.dumps(example, ensure_ascii=False, indent=2), encoding="utf-8")

def money(m):
    return f'{m["amount"]} {m["currency"]}' if m else ""

# Flat rows: one per service x weight tier (the shape a shop needs for a rate table)
tier_cols = ["service_id", "mode", "destination", "carrier", "carrier_ids", "variant", "max_kg", "tier_variant",
             "price_depot", "price_zpoint", "currency", "lead_days_min", "lead_days_max", "cod_max", "insurance_max", "suspended"]
tier_rows = []
for s in picked:
    for t in s["tiers"]:
        tier_rows.append([s["id"], s["mode"], s["destination"], s["carrierName"], " ".join(map(str, s["carrierIds"])),
                          s["variant"] or "", t["maxKg"], t["variant"] or "", t["price"].get("depot"), t["price"].get("zpoint"),
                          s["currency"], (s["leadTimeDays"] or {}).get("min"), (s["leadTimeDays"] or {}).get("max"),
                          money(s["cod"]["max"]) if s["cod"]["available"] else "not offered", money(s["insurance"]["max"]),
                          "yes" if s["suspended"] else "no"])

with open(out / "sk-example-tiers.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")  # ";" + BOM so Excel with SK/CZ locale opens it correctly
    w.writerow(tier_cols)
    w.writerows(tier_rows)

# Excel workbook
wb = Workbook()
head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="BA0C2F")

def sheet(title, cols, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font, c.fill, c.alignment = head_font, head_fill, Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(cols, 1):
        width = max(len(str(col)), *(len(str(r[i - 1])) for r in rows)) if rows else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    return ws

src_info = pl["source"]
sheet("About", ["field", "value"], [
    ["source", src_info.get("id")], ["publisher", src_info.get("publisher", "Packeta Slovakia s.r.o.")],
    ["valid from", src_info.get("validFrom")], ["language", src_info.get("language")],
    ["prices", "EUR, excl. VAT; depot = handed in at a Packeta depot, zpoint = handed in at a Z-POINT"],
    ["note", "EXAMPLE: 3-5 services per delivery mode. Unofficial; always verify against the source PDF."],
], first=True)

ws = sheet("Price tiers", tier_cols, tier_rows)
for row in ws.iter_rows(min_row=2, min_col=9, max_col=10):
    for c in row:
        c.number_format = "0.00"

sheet("Services", ["service_id", "mode", "destination", "carrier", "title", "lead_days", "max_weight_kg", "max_side_cm", "max_sum_cm",
                   "cod_available", "cod_fee", "card_payment_pct", "volumetric_divisor", "page"],
      [[s["id"], s["mode"], s["destination"], s["carrierName"], s["title"],
        f'{s["leadTimeDays"]["min"]}-{s["leadTimeDays"]["max"]}' if s["leadTimeDays"] else "",
        s["limits"].get("maxWeightKg"), s["limits"].get("maxSideCm"), s["limits"].get("maxSumCm"),
        "yes" if s["cod"]["available"] else "no", "; ".join(f'up to {money(t["upTo"])}: {t.get("fee", "")}' for t in s["cod"]["tiers"]),
        s["cod"].get("cardPaymentPct"), (s["volumetric"] or {}).get("divisor"), s["page"]] for s in picked])

sheet("Insurance", ["service_id", "up_to", "fee_eur", "fee_pct", "included"],
      [[s["id"], money(t["upTo"]), t.get("fee"), t.get("feePct"), "yes" if t.get("included") else ""]
       for s in picked for t in s["insurance"]["tiers"]])

sheet("Service fees", ["service_id", "code", "label", "amount", "currency", "by_weight", "note"],
      [[s["id"], f["code"], f["label"], f["amount"], f["currency"],
        "; ".join(f'≤{b["maxKg"]} kg: {b["amount"]}' for b in (f["byWeight"] or [])), f["note"] or ""]
       for s in picked for f in s["fees"]])

sheet("Returns", ["from_country", "price", "currency"], [[r["fromCountry"], r["price"], r["currency"]] for r in pl["returns"]["prices"]])

sur = pl["surcharges"]
sheet("Surcharges", ["item", "from_eur_per_1000l", "to_eur_per_1000l", "pct_or_amount"],
      [["toll per started kg", "", "", sur["toll"]["perStartedKg"] if sur["toll"] else ""]] +
      [["fuel (diesel price band)", d["fromEurPer1000l"], d["toEurPer1000l"], d["pct"]] for d in sur["fuel"]["dieselTable"]])

sheet("General fees", ["code", "label", "amount", "currency", "unit"], [[f["code"], f["label"], f["amount"], f["currency"], f["unit"] or ""] for f in pl["fees"]])
sheet("Penalties", ["code", "label", "amount", "currency", "unit"], [[f["code"], f["label"], f["amount"], f["currency"], f["unit"] or ""] for f in pl["penalties"]])

wb.save(out / "sk-example.xlsx")
print(f"{len(picked)} services, {len(tier_rows)} tier rows ->", out)
